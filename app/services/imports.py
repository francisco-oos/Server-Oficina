from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import unicodedata
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import load_settings
from app.db.models import AttendanceRecord, EmploymentEngagement, ImportBatch, ImportIssue, Person, WorkGroup, GroupAssignment
from app.services.events import record_event


ALIASES = {
    "employment_id": {"CLAVE", "ID", "ID EMPLEADO", "NO EMPLEADO", "NUM EMPLEADO", "NUMERO EMPLEADO"},
    "name": {"NOMBRE", "NOMBRE COMPLETO", "EMPLEADO", "TRABAJADOR"},
    "position": {"PUESTO", "CATEGORIA", "CATEGORÍA"},
    "group": {"GRUPO", "CUADRILLA", "BRIGADA"},
    "status": {"ESTADO", "ESTATUS", "STATUS"},
}


def normalize_text(value: object) -> str:
    text = "" if value is None else str(value).strip()
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", text).strip().upper()


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _read_rows(filename: str, content: bytes) -> list[dict[str, object]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        text = content.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        return [dict(r) for r in reader]
    if suffix in {".xlsx", ".xlsm"}:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(x).strip() if x is not None else "" for x in rows[0]]
        return [{headers[i]: row[i] if i < len(row) else None for i in range(len(headers))} for row in rows[1:] if any(v is not None for v in row)]
    raise ValueError("Formato no soportado; use .csv, .xlsx o .xlsm")


def _canonical_headers(row: dict[str, object]) -> dict[str, str]:
    headers = {normalize_text(k): k for k in row.keys()}
    result: dict[str, str] = {}
    for target, options in ALIASES.items():
        for option in options:
            if normalize_text(option) in headers:
                result[target] = headers[normalize_text(option)]
                break
    return result


def _parse_date_header(value: str) -> date | None:
    raw = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y", "%d-%m-%y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    return None


def _resolve_person(db: Session, employment_id: str | None, name: str | None) -> tuple[Person | None, str | None]:
    id_not_found = False
    if employment_id:
        eng = db.scalar(select(EmploymentEngagement).where(EmploymentEngagement.employment_id == employment_id).order_by(EmploymentEngagement.start_date.desc()))
        if eng:
            return db.get(Person, eng.person_id), None
        id_not_found = True
    if name:
        normalized = normalize_text(name)
        matches = db.scalars(select(Person).where(Person.normalized_name == normalized)).all()
        if len(matches) == 1:
            # Si el archivo trae un ID laboral desconocido pero el nombre coincide con una persona existente,
            # puede ser recontratación/cambio de ID. No se decide automáticamente.
            return matches[0], "POSSIBLE_ID_CHANGE" if id_not_found else None
        if len(matches) > 1:
            return None, "AMBIGUOUS_NAME"
    return None, "NOT_FOUND"


def preview_import(db: Session, *, kind: str, filename: str, content: bytes, user_id: str) -> ImportBatch:
    settings = load_settings()
    rows = _read_rows(filename, content)
    batch = ImportBatch(kind=kind, original_name=filename, stored_path="", sha256=sha256_bytes(content), created_by=user_id)
    db.add(batch)
    db.flush()
    folder = settings.data_dir / "imports" / batch.id
    folder.mkdir(parents=True, exist_ok=True)
    original_path = folder / ("original" + Path(filename).suffix.lower())
    original_path.write_bytes(content)
    batch.stored_path = str(original_path)

    summary = {"rows": len(rows), "matched": 0, "new": 0, "review": 0, "issues": 0, "attendance_cells": 0}
    staged: list[dict] = []
    if rows:
        columns = _canonical_headers(rows[0])
        if "name" not in columns and "employment_id" not in columns:
            raise ValueError("No se encontró columna NOMBRE ni CLAVE/ID EMPLEADO")
        for idx, row in enumerate(rows, start=2):
            employment_id = str(row.get(columns.get("employment_id", ""), "") or "").strip() or None
            name = str(row.get(columns.get("name", ""), "") or "").strip() or None
            person, issue = _resolve_person(db, employment_id, name)
            if person and not issue:
                summary["matched"] += 1
            elif person and issue:
                summary["review"] += 1
            else:
                summary["new"] += 1
            if issue in {"AMBIGUOUS_NAME", "POSSIBLE_ID_CHANGE"}:
                message = f"Nombre ambiguo: {name}" if issue == "AMBIGUOUS_NAME" else f"Posible recontratación o cambio de ID laboral: {name} / {employment_id}"
                db.add(ImportIssue(batch_id=batch.id, row_number=idx, issue_type=issue, message=message, payload={"employment_id": employment_id, "name": name, "person_id": person.id if person else None}))
                summary["issues"] += 1
            attendance = []
            if kind == "attendance":
                for header, value in row.items():
                    d = _parse_date_header(header)
                    if d and value not in (None, ""):
                        attendance.append({"date": d.isoformat(), "status": str(value).strip()})
                summary["attendance_cells"] += len(attendance)
            staged.append({
                "row_number": idx,
                "person_id": person.id if person else None,
                "employment_id": employment_id,
                "name": name,
                "position": str(row.get(columns.get("position", ""), "") or "").strip() or None,
                "group": str(row.get(columns.get("group", ""), "") or "").strip() or None,
                "status": str(row.get(columns.get("status", ""), "") or "").strip() or None,
                "attendance": attendance,
                "blocked": issue in {"AMBIGUOUS_NAME", "POSSIBLE_ID_CHANGE"},
            })
    staged_path = folder / "staged.json"
    staged_path.write_text(json.dumps(staged, ensure_ascii=False, indent=2), encoding="utf-8")
    summary["staged_path"] = str(staged_path)
    batch.summary = summary
    db.commit()
    db.refresh(batch)
    return batch


def commit_import(db: Session, batch: ImportBatch, user_id: str) -> dict:
    if batch.status != "PREVIEW":
        raise ValueError("El lote ya fue confirmado o no está en PREVIEW")
    staged_path = Path(batch.summary["staged_path"])
    rows = json.loads(staged_path.read_text(encoding="utf-8"))
    committed = 0
    attendance_count = 0
    for item in rows:
        if item.get("blocked"):
            continue
        person = db.get(Person, item.get("person_id")) if item.get("person_id") else None
        if not person:
            if not item.get("name"):
                continue
            person = Person(full_name=item["name"], normalized_name=normalize_text(item["name"]))
            db.add(person)
            db.flush()
            record_event(db, entity_type="PERSON", entity_id=person.id, event_type="PERSON_CREATED", source_type="IMPORT", source_id=batch.id, actor_user_id=user_id)
            if item.get("employment_id"):
                eng = EmploymentEngagement(
                    person_id=person.id,
                    employment_id=item["employment_id"],
                    employer_type="DIRECT",
                    position=item.get("position"),
                    start_date=date.today(),
                    status=item.get("status") or "ACTIVE",
                )
                db.add(eng)
        if item.get("group"):
            code = normalize_text(item["group"])
            group = db.scalar(select(WorkGroup).where(WorkGroup.project_id.is_(None), WorkGroup.code == code))
            if not group:
                group = WorkGroup(project_id=None, code=code, name=item["group"])
                db.add(group); db.flush()
            active = db.scalar(select(GroupAssignment).where(GroupAssignment.person_id == person.id, GroupAssignment.end_date.is_(None)))
            if not active or active.group_id != group.id:
                if active:
                    active.end_date = date.today()
                db.add(GroupAssignment(person_id=person.id, group_id=group.id, start_date=date.today()))
        for att in item.get("attendance", []):
            d = date.fromisoformat(att["date"])
            rec = db.scalar(select(AttendanceRecord).where(AttendanceRecord.person_id == person.id, AttendanceRecord.attendance_date == d))
            if rec:
                rec.status = att["status"]
                rec.import_batch_id = batch.id
            else:
                db.add(AttendanceRecord(person_id=person.id, attendance_date=d, status=att["status"], import_batch_id=batch.id))
            attendance_count += 1
        committed += 1
    batch.status = "COMMITTED"
    batch.committed_at = datetime.now(timezone.utc)
    db.commit()
    return {"rows_committed": committed, "attendance_records": attendance_count, "issues_open": len(db.scalars(select(ImportIssue).where(ImportIssue.batch_id == batch.id, ImportIssue.status == "OPEN")).all())}
